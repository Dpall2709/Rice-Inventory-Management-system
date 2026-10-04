"""
The truck register: one row per sale (truck), mill to party, the way rice
traders keep it in their Excel sheet -

    date · truck · mill · purchase rates and bags · weight · loading ·
    freight advance + balance · total cost · sale rate · unloading date ·
    party weight · shortage · party value · broker · CD · brokerage ·
    party owes · received · last payment date · profit / loss · status

Two things the hand-made sheets usually get wrong are done properly here:

  * freight the party paid the driver is counted ONCE - as your cost, and as
    a cut from what the party pays - not twice;
  * a truck that is not paid yet is "awaiting payment", not a loss of its
    whole value.
"""

from collections import OrderedDict
from decimal import Decimal

from django.db.models import Max
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy

from core.models import Payment, Sale
from core.services.customer_ledger import status_map_for_sales
from core.services.sale_service import earned_profit, money, payment_status, sale_profit

ON_THE_WAY = "on_the_way"
AWAITING = "awaiting_payment"
PARTIAL = "partly_paid"
SETTLED = "settled"

STATUS_LABELS = {
    ON_THE_WAY: gettext_lazy("On the way"),
    AWAITING: gettext_lazy("Awaiting payment"),
    PARTIAL: gettext_lazy("Partly paid"),
    SETTLED: gettext_lazy("Settled"),
}


def _purchase_lines(lots):
    """[(mill name, rate per kg, bags)] grouped like the sheet's Rate(1)/Packet(1)…"""
    grouped = OrderedDict()
    for lot in lots:
        key = (lot.mill.mill_name, Decimal(lot.buy_rate_per_kg))
        grouped[key] = grouped.get(key, 0) + int(lot.bag_count)
    return [(mill, rate, bags) for (mill, rate), bags in grouped.items()]


def register_rows(company, sales):
    """Build the register for an iterable of this company's sales (oldest first)."""
    sales = list(
        sales.select_related("customer", "broker")
        .prefetch_related("lots__mill", "lots__purchase_item__purchase")
        .order_by("sale_date", "id")
    )
    status_map = status_map_for_sales(company, sales)

    last_paid = {
        row["sale"]: row["last"]
        for row in Payment.objects.filter(company=company, related_type="sale", sale__in=sales)
        .values("sale").annotate(last=Max("payment_date"))
    }

    rows = []
    for sale in sales:
        profit = sale_profit(sale)
        figures = profit["settlement"]
        ledger = status_map.get(sale.id, {})
        paid = ledger.get("paid", Decimal("0"))
        due = ledger.get("due", figures["net_receivable"])

        if not sale.unload_date:
            status = ON_THE_WAY
        elif due <= 0:
            status = SETTLED
        elif paid > 0:
            status = PARTIAL
        else:
            status = AWAITING

        kg = Decimal(sale.total_quantity_kg or 0)
        earned, pending = earned_profit(profit["profit"], figures["net_receivable"], paid)
        rows.append({
            "sale": sale,
            "mills": sorted({lot.mill.mill_name for lot in profit["lots"]}),
            "purchase_lines": _purchase_lines(profit["lots"]),
            "weight_kg": money(kg),
            "bags": sale.total_bags,
            "goods_cost": profit["goods_cost"],
            "loading": profit["loading"],
            "freight_advance": figures["freight_paid_by_us"],
            "freight_balance": figures["freight_paid_by_party"],
            "freight_cost": figures["freight_cost"],
            "total_cost": profit["cost"],
            "sale_rate": money(Decimal(sale.taxable_amount or 0) / kg) if kg else Decimal("0.00"),
            "unload_date": sale.unload_date,
            "received_kg": figures["received_kg"] if figures["unloaded"] else None,
            "loss_kg": figures["weight_loss_kg"] if figures["unloaded"] else None,
            "party_value": figures["party_value"],
            "cash_discount": figures["cash_discount"],
            "commission": figures["commission"],
            "brokerage_cut": figures["brokerage_cut"],
            "net_receivable": figures["net_receivable"],
            "received": money(paid),
            "due": money(due),
            "last_payment": last_paid.get(sale.id) or (sale.sale_date if sale.advance_received else None),
            "profit": profit["profit"],            # total profit when fully paid
            "profit_earned": earned,               # in hand, by money received
            "profit_pending": pending,             # comes in with the rest of the money
            "payment_status": payment_status(paid, due),
            "paid_percent": min(int(paid * 100 / figures["net_receivable"]), 100) if figures["net_receivable"] > 0 else 100,
            "status": status,
            "status_label": STATUS_LABELS[status],
        })
    return rows


PAYMENT_LABELS = {
    "paid": gettext_lazy("Paid"),
    "partial": gettext_lazy("Partially paid"),
    "due": gettext_lazy("Due"),
}


def register_totals(rows):
    def total(key):
        return money(sum((r[key] or Decimal("0") for r in rows), Decimal("0")))

    settled = [r for r in rows if r["status"] == SETTLED and r["profit"] is not None]
    open_rows = [r for r in rows if r["status"] != SETTLED and r["profit"] is not None]
    costed = [r for r in rows if r["profit"] is not None]
    return {
        "sale_amount": total("net_receivable"),
        "cost": money(sum((r["total_cost"] for r in costed), Decimal("0"))),
        "profit_total": money(sum((r["profit"] for r in costed), Decimal("0"))),
        "profit_earned": money(sum((r["profit_earned"] for r in costed), Decimal("0"))),
        "profit_pending": money(sum((r["profit_pending"] for r in costed), Decimal("0"))),
        "loss_total": money(-sum((r["profit"] for r in costed if r["profit"] < 0), Decimal("0"))),
        "bills_paid": sum(1 for r in rows if r["payment_status"] == "paid"),
        "bills_partial": sum(1 for r in rows if r["payment_status"] == "partial"),
        "bills_due": sum(1 for r in rows if r["payment_status"] == "due"),
        "trucks": len(rows),
        "bags": sum(int(r["bags"] or 0) for r in rows),
        "weight_kg": total("weight_kg"),
        "investment": total("total_cost"),
        "party_owes": total("net_receivable"),
        "received": total("received"),
        "outstanding": total("due"),
        "profit_settled": money(sum((r["profit"] for r in settled), Decimal("0"))),
        "profit_expected": money(sum((r["profit"] for r in open_rows), Decimal("0"))),
        "loss_trucks": sum(1 for r in rows if r["profit"] is not None and r["profit"] < 0),
        "no_cost": sum(1 for r in rows if r["profit"] is None),
        "on_the_way": sum(1 for r in rows if r["status"] == ON_THE_WAY),
    }


def group_summary(rows, key):
    """
    Profit and money per broker or per customer, for comparing who is worth
    dealing with. `key` is "broker" or "customer".
    """
    groups = OrderedDict()
    for r in rows:
        sale = r["sale"]
        if key == "broker":
            ident = sale.broker_id or 0
            name = sale.broker.broker_name if sale.broker else str(_("Direct (no broker)"))
        else:
            ident = sale.customer_id or ("name", sale.customer_name)
            name = sale.customer_name
        group = groups.setdefault(ident, {
            "id": ident if isinstance(ident, int) else None, "name": name, "trucks": 0,
            "weight_kg": Decimal("0"), "loss_kg": Decimal("0"), "party_owes": Decimal("0"),
            "received": Decimal("0"), "due": Decimal("0"), "cash_discount": Decimal("0"),
            "brokerage": Decimal("0"), "profit": Decimal("0"), "losses": 0,
            "cost": Decimal("0"), "profit_earned": Decimal("0"), "profit_pending": Decimal("0"),
        })
        group["trucks"] += 1
        group["weight_kg"] += r["weight_kg"] or 0
        group["loss_kg"] += r["loss_kg"] or 0
        group["party_owes"] += r["net_receivable"] or 0
        group["received"] += r["received"] or 0
        group["due"] += r["due"] or 0
        group["cash_discount"] += r["cash_discount"] or 0
        group["brokerage"] += r["commission"] if sale.broker_id else 0
        if r["profit"] is not None:
            group["profit"] += r["profit"]
            group["cost"] += r["total_cost"]
            group["profit_earned"] += r["profit_earned"]
            group["profit_pending"] += r["profit_pending"]
            if r["profit"] < 0:
                group["losses"] += 1
    result = list(groups.values())
    for group in result:
        for field in ("weight_kg", "loss_kg", "party_owes", "received", "due", "cash_discount", "brokerage",
                      "profit", "cost", "profit_earned", "profit_pending"):
            group[field] = money(group[field])
    result.sort(key=lambda g: g["profit"], reverse=True)
    return result


def register_excel(company, rows, totals, title):
    """The register as an .xlsx workbook (HttpResponse)."""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = _("Truck register")[:31]

    ws.append([company.company_name])
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])

    headers = [
        _("S. No"), _("Date"), _("Invoice"), _("Truck no"), _("Mill"), _("Purchase rate × bags"),
        _("Bags"), _("Weight (kg)"), _("Rice cost"), _("Loading"), _("Freight advance"),
        _("Freight balance (party)"), _("Total cost"), _("Sale rate/kg"), _("Customer"),
        _("Unloaded on"), _("Party weight (kg)"), _("Shortage (kg)"), _("Party value"), _("Broker"),
        _("Who pays us"),
        _("Cash discount"), _("Brokerage"), _("Party owes"), _("Received"), _("Last payment"),
        _("Due"), _("Full profit / loss"), _("Profit earned"), _("Payment"), _("Status"),
    ]
    ws.append(headers)
    header_row = ws.max_row
    fill = PatternFill("solid", fgColor="1F4E3D")
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    def num(value):
        return float(value) if value is not None else None

    for index, r in enumerate(rows, start=1):
        sale = r["sale"]
        ws.append([
            index, sale.sale_date, sale.invoice_no, sale.vehicle_number, ", ".join(r["mills"]),
            "; ".join(f"{rate} × {bags}" for _mill, rate, bags in r["purchase_lines"]),
            r["bags"], num(r["weight_kg"]), num(r["goods_cost"]), num(r["loading"]),
            num(r["freight_advance"]), num(r["freight_balance"]), num(r["total_cost"]),
            num(r["sale_rate"]), sale.customer_name, r["unload_date"], num(r["received_kg"]),
            num(r["loss_kg"]), num(r["party_value"]), sale.broker.broker_name if sale.broker else "",
            str(sale.get_collect_from_display()) if sale.broker_id else str(_("Customer pays us")),
            num(r["cash_discount"]), num(r["commission"]), num(r["net_receivable"]), num(r["received"]),
            r["last_payment"], num(r["due"]), num(r["profit"]), num(r["profit_earned"]),
            str(PAYMENT_LABELS[r["payment_status"]]), str(r["status_label"]),
        ])
        if r["profit"] is not None:
            for column in (28, 29):
                ws.cell(row=ws.max_row, column=column).font = Font(
                    color="B91C1C" if r["profit"] < 0 else "15803D", bold=True)

    ws.append([])
    for label, value in [
        (_("Trucks"), totals["trucks"]),
        (_("Paid"), totals["bills_paid"]),
        (_("Part paid"), totals["bills_partial"]),
        (_("Unpaid"), totals["bills_due"]),
        (_("Total sale"), num(totals["sale_amount"])),
        (_("Total cost"), num(totals["investment"])),
        (_("Received"), num(totals["received"])),
        (_("Still to receive"), num(totals["outstanding"])),
        (_("Profit earned (on money received)"), num(totals["profit_earned"])),
        (_("Profit still to come with the dues"), num(totals["profit_pending"])),
        (_("Full profit when all is paid"), num(totals["profit_total"])),
        (_("Loss"), num(totals["loss_total"])),
    ]:
        ws.append([label, None, None, None, value])  # label spills over the empty cells
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    for column in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(column)].width = 14
    ws.column_dimensions["E"].width = 26
    ws.column_dimensions["F"].width = 24
    ws.column_dimensions["O"].width = 26
    ws.freeze_panes = ws.cell(row=header_row + 1, column=5)

    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
            elif hasattr(cell.value, "year"):
                cell.number_format = "DD-MMM-YYYY"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="truck-register.xlsx"'
    wb.save(response)
    return response
