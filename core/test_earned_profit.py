"""
Profit counts as earned only as the money comes in, and every screen
(truck register, broker / customer summaries, sale bills, dashboard) uses the
same rule, so their numbers never disagree.
"""

import io
from datetime import timedelta
from decimal import Decimal

from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.models import Payment, Sale
from core.services.dashboard import period_kpis
from core.services.party_reports import parse_period
from core.services.sale_service import earned_profit, payment_status
from core.services.trade_register import group_summary, register_rows, register_totals
from core import test_settlement

TRUCK_PROFIT = Decimal("18852.70")
PARTY_OWES = Decimal("954852.70")


class TruckFixture(TestCase):
    """The settlement tests' truck set-up, without re-running their tests."""
    setUp = test_settlement.TruckSettlementTests.setUp
    truck_post = test_settlement.TruckSettlementTests.truck_post
    make_truck = test_settlement.TruckSettlementTests.make_truck
    settle = test_settlement.TruckSettlementTests.settle


class EarnedProfitRuleTests(SimpleTestCase):

    def test_share_of_the_money_received(self):
        self.assertEqual(earned_profit(Decimal("1000"), Decimal("10000"), 0), (Decimal("0.00"), Decimal("1000.00")))
        self.assertEqual(earned_profit(Decimal("1000"), Decimal("10000"), 2500), (Decimal("250.00"), Decimal("750.00")))
        self.assertEqual(earned_profit(Decimal("1000"), Decimal("10000"), 10000), (Decimal("1000.00"), Decimal("0.00")))

    def test_overpayment_never_earns_more_than_the_full_profit(self):
        self.assertEqual(earned_profit(Decimal("1000"), Decimal("10000"), 12000)[0], Decimal("1000.00"))

    def test_a_loss_counts_in_full_at_once(self):
        self.assertEqual(earned_profit(Decimal("-500"), Decimal("10000"), 0), (Decimal("-500.00"), Decimal("0.00")))

    def test_unknown_cost(self):
        self.assertEqual(earned_profit(None, Decimal("10000"), 0), (None, None))

    def test_payment_status(self):
        self.assertEqual(payment_status(Decimal("0"), Decimal("10")), "due")
        self.assertEqual(payment_status(Decimal("5"), Decimal("5")), "partial")
        self.assertEqual(payment_status(Decimal("10"), Decimal("0")), "paid")


class EarnedProfitEverywhereTests(TruckFixture):

    def pay(self, sale, amount):
        response = self.client.post(reverse("add_sale_payment", args=[sale.id]), {
            "amount": str(amount), "payment_mode": "Bank", "payment_date": timezone.localdate().isoformat(),
        })
        self.assertEqual(response.status_code, 302)

    def figures(self):
        rows = register_rows(self.company, Sale.objects.for_company(self.company))
        return rows, register_totals(rows)

    def test_profit_follows_the_payments(self):
        sale = self.settle(self.make_truck())

        rows, totals = self.figures()
        self.assertEqual(rows[0]["payment_status"], "due")
        self.assertEqual(totals["profit_earned"], Decimal("0.00"))
        self.assertEqual(totals["profit_total"], TRUCK_PROFIT)
        self.assertEqual(totals["bills_due"], 1)

        self.pay(sale, PARTY_OWES / 2)
        rows, totals = self.figures()
        self.assertEqual(rows[0]["payment_status"], "partial")
        self.assertEqual(totals["profit_earned"], Decimal("9426.35"))
        self.assertEqual(totals["profit_pending"], Decimal("9426.35"))
        self.assertEqual(totals["bills_partial"], 1)

        self.pay(sale, PARTY_OWES / 2)
        rows, totals = self.figures()
        self.assertEqual(rows[0]["payment_status"], "paid")
        self.assertEqual(totals["profit_earned"], TRUCK_PROFIT)
        self.assertEqual(totals["outstanding"], Decimal("0.00"))
        self.assertEqual(totals["bills_paid"], 1)

        # deleting the payment takes the earned profit back
        payment = Payment.objects.filter(sale=sale).order_by("-id").first()
        self.client.post(reverse("delete_payment", args=[payment.id]))
        self.assertEqual(self.figures()[1]["profit_earned"], Decimal("9426.35"))

    def test_every_screen_shows_the_same_numbers(self):
        sale = self.settle(self.make_truck())
        self.pay(sale, PARTY_OWES / 2)
        rows, totals = self.figures()

        broker = group_summary(rows, "broker")[0]
        customer = group_summary(rows, "customer")[0]
        for group in (broker, customer):
            self.assertEqual(group["profit_earned"], totals["profit_earned"])
            self.assertEqual(group["received"], totals["received"])
            self.assertEqual(group["due"], totals["outstanding"])

        today = timezone.localdate()
        kpis = period_kpis(self.company, today.replace(day=1, month=1), today)
        self.assertEqual(kpis["profit_earned"], totals["profit_earned"])
        self.assertEqual(kpis["profit"], TRUCK_PROFIT)
        self.assertEqual(kpis["bill_due"], totals["outstanding"])
        self.assertEqual(kpis["bills_partial"], 1)

        detail = self.client.get(reverse("sale_detail", args=[sale.id]))
        self.assertEqual(detail.context["profit"]["earned"], Decimal("9426.35"))
        self.assertEqual(detail.context["profit"]["payment_status"], "partial")

        listing = self.client.get(reverse("sale_list"))
        self.assertEqual(listing.context["total_profit_earned"], Decimal("9426.35"))

        for name in ("trade_register", "dashboard"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 200)
        self.assertEqual(self.client.get(reverse("broker_report_detail", args=[self.broker.id])).context["profit_earned"],
                         Decimal("9426.35"))
        self.assertEqual(self.client.get(reverse("customer_ledger", args=[self.customer.id])).context["profit_earned"],
                         Decimal("9426.35"))

    def test_register_excel_has_the_trucks_and_earned_profit(self):
        sale = self.settle(self.make_truck())
        self.pay(sale, PARTY_OWES / 2)
        response = self.client.get(reverse("trade_register") + "?export=xlsx")
        sheet = load_workbook(io.BytesIO(response.content)).active
        values = [cell for row in sheet.iter_rows(values_only=True) for cell in row]
        self.assertIn("CG04PX-5695", values)
        self.assertIn(9426.35, values)
        self.assertIn("Partially paid", values)


class StatementDownloadTests(TruckFixture):

    def test_no_dates_means_the_whole_account(self):
        start, end, _label = parse_period({})
        self.assertEqual(end, timezone.localdate())
        self.assertLess(start.year, 2001)

    def test_old_trucks_are_in_a_download_without_dates(self):
        sale = self.make_truck(sale_date=(timezone.localdate().replace(day=1) - timedelta(days=40)).isoformat())
        self.settle(sale)
        response = self.client.get(reverse("broker_statement_export", args=[self.broker.id, "xlsx"]))
        book = load_workbook(io.BytesIO(response.content))
        values = [cell for ws in book for row in ws.iter_rows(values_only=True) for cell in row]
        self.assertIn(sale.invoice_no, values)

    def test_customer_sheet_lists_trucks_the_broker_pays_for(self):
        sale = self.settle(self.make_truck())
        response = self.client.get(reverse("customer_statement_export", args=[self.customer.id, "xlsx"]))
        account = load_workbook(io.BytesIO(response.content)).active
        text = " ".join(str(cell) for row in account.iter_rows(values_only=True) for cell in row if cell)
        self.assertIn(sale.invoice_no, text)
